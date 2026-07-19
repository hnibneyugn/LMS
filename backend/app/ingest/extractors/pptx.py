"""PPTX: one slide becomes one '## <title>' section so the splitter can cut it."""

import io

from pptx import Presentation

from app.ingest.extractors.errors import ExtractError


def extract(data: bytes) -> str:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as err:
        raise ExtractError(
            "Không mở được file .pptx — file có thể hỏng hoặc không đúng định dạng."
        ) from err

    parts: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = _title_of(slide) or f"Slide {number}"
        parts.append(f"## {title}")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text and text != title:
                parts.append(text)
    return "\n\n".join(parts).strip()


def _title_of(slide) -> str:
    placeholder = slide.shapes.title
    if placeholder is None:
        return ""
    return placeholder.text.strip()
