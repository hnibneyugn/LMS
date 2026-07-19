import io

import pytest
from docx import Document
from pptx import Presentation

from app.ingest.extractors import EXTRACTORS, ExtractError
from app.ingest.extractors import md as md_extractor
from app.ingest.extractors import pdf as pdf_extractor


def test_registry_covers_all_four_formats():
    assert set(EXTRACTORS) == {"md", "docx", "pptx", "pdf"}


def test_md_keeps_the_body_and_drops_frontmatter():
    raw = b"---\ntitle: Bai 1\n---\n\n# Chuong 1\nnoi dung\n"
    result = md_extractor.extract(raw)
    assert "# Chuong 1" in result
    assert "title: Bai 1" not in result


def test_md_without_frontmatter_is_unchanged():
    raw = b"# Chuong 1\nnoi dung\n"
    assert md_extractor.extract(raw).strip() == "# Chuong 1\nnoi dung"


def test_docx_maps_headings_and_lists_to_markdown():
    document = Document()
    document.add_heading("Chuong mot", level=1)
    document.add_paragraph("doan van thuong")
    document.add_paragraph("mot muc", style="List Bullet")
    document.add_heading("Muc con", level=2)
    buffer = io.BytesIO()
    document.save(buffer)

    result = EXTRACTORS["docx"](buffer.getvalue())
    assert "# Chuong mot" in result
    assert "## Muc con" in result
    assert "- mot muc" in result
    assert "doan van thuong" in result


def test_pptx_turns_each_slide_into_an_h2():
    presentation = Presentation()
    for title_text in ("Slide mot", "Slide hai"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title_text
        slide.placeholders[1].text = f"noi dung {title_text}"
    buffer = io.BytesIO()
    presentation.save(buffer)

    result = EXTRACTORS["pptx"](buffer.getvalue())
    assert "## Slide mot" in result
    assert "## Slide hai" in result
    assert "noi dung Slide hai" in result


def test_pdf_without_text_layer_raises_a_vietnamese_error():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ExtractError) as err:
        pdf_extractor.extract(buffer.getvalue())
    assert "không có văn bản" in str(err.value)


def test_corrupt_bytes_raise_extract_error_not_a_library_exception():
    for file_type in ("docx", "pptx", "pdf"):
        with pytest.raises(ExtractError):
            EXTRACTORS[file_type](b"day khong phai file hop le")
